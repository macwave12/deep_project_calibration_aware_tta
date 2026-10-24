"""The deck is built from the committed CSVs and records; these tests pin that."""

import json
from pathlib import Path

import pandas as pd
import pytest
from pptx import Presentation

from presentation.build_pptx import build

ROOT = Path(__file__).resolve().parents[1]


def _slide_text(slide):
    parts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            parts.append(shape.text_frame.text)
        if shape.has_table:
            for row in shape.table.rows:
                parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


@pytest.fixture(scope="module")
def deck(tmp_path_factory):
    out_dir = tmp_path_factory.mktemp("deck")
    out = build(
        ROOT / "results/summary.csv",
        ROOT / "results/ablation.csv",
        ROOT / "results/ablation_gate.csv",
        ROOT / "results/raw/ablation",
        ROOT / "configs/method/ablation/margin.yaml",
        ROOT / "configs/method/cal_tda_gate.yaml",
        out_dir / "deck.pptx",
        markdown_path=out_dir / "slides.md",
    )
    prs = Presentation(str(out))
    return prs, [_slide_text(s) for s in prs.slides], (out_dir / "slides.md").read_text(encoding="utf-8")


def test_deck_has_eighteen_slides_each_with_speaker_notes(deck):
    prs, _, markdown = deck
    assert len(prs.slides) == 18
    for index, slide in enumerate(prs.slides, start=1):
        assert slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip(), index
    assert markdown.count("\n## ") == 18


def test_slides_name_no_code_identifiers_and_no_old_story(deck):
    _, texts, _ = deck
    banned = ("cal_tda", "clipzs", "Limitations", "Ours, all 3", "second ablation")
    for index, text in enumerate(texts, start=1):
        for token in banned:
            assert token not in text, (index, token)


def test_headline_slide_quotes_the_summary_csv_means(deck):
    _, texts, _ = deck
    summary = pd.read_csv(ROOT / "results/summary.csv")
    means = summary.groupby("method")[["accuracy", "ece"]].mean().round(2)
    slide = next(t for t in texts if "Ours: gate + temp" in t)
    for method in ("clipzs", "tda", "cal_tda_gate"):
        assert f"{means.loc[method, 'accuracy']:.2f}" in slide, method
        assert f"{means.loc[method, 'ece']:.2f}" in slide, method


def test_gate_result_slide_matches_ablation_gate_csv(deck):
    _, texts, _ = deck
    gate = pd.read_csv(ROOT / "results/ablation_gate.csv")
    means = gate.groupby("variant")[["accuracy", "ece"]].mean().round(2)
    slide = next(t for t in texts if "gate ECE" in t)
    assert f"{means.loc['gate', 'ece']:.2f}" in slide
    assert f"{means.loc['none', 'ece']:.2f}" in slide
    assert f"{means.loc['gate', 'accuracy']:.2f}" in slide


def test_margin_slide_shows_admission_rates_from_the_raw_records(deck):
    _, texts, _ = deck
    slide = next(t for t in texts if "margin ECE" in t)
    for dataset in ("aircraft", "pets"):
        record = json.loads((ROOT / f"results/raw/ablation/{dataset}_margin.json").read_text(encoding="utf-8"))
        rate = 100.0 * record["hyperparams"]["admission"]["admission_rate"]
        assert f"{rate:.1f}%" in slide, dataset
    # the dataset table folded into this slide
    assert "1,692" in slide and "8,100" in slide


def test_ablation_slide_lists_the_four_margin_off_variants(deck):
    _, texts, _ = deck
    gate = pd.read_csv(ROOT / "results/ablation_gate.csv")
    means = gate.groupby("variant")[["accuracy", "ece"]].mean().round(2)
    slide = next(t for t in texts if "+ gated fusion + temperature" in t)
    for variant in ("none", "gate", "temp", "gate_temp"):
        assert f"{means.loc[variant, 'ece']:.2f}" in slide, variant
